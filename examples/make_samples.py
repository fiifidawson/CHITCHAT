"""Build the binary sample documents in examples/my_docs/: .docx, .pdf and .xlsx.

The text samples (.md, .txt, .eml, .html) are written by hand. Run from the repo root:

    uv run python examples/make_samples.py

Everything is fictional: names, numbers, addresses and IDs are made up.
"""

from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).parent / "my_docs"
NOTE = "Fictional sample for testing. All names, numbers, addresses and IDs are made up."

CONSENT_FORM = [
    ("title", "Consent form: use of helpline calls to improve the chatbot"),
    ("note", NOTE),
    ("p", "Organisation: Northbridge Relief Network (fictional). Version 2, 15 March 2026."),
    ("h", "What we are asking"),
    (
        "p",
        "We would like to use a written copy of this call, with your name, phone number, "
        "camp block and case number removed, to teach our helpline chatbot to give better answers.",
    ),
    ("h", "What happens if you say yes"),
    (
        "p",
        "A staff member removes your personal details. The text is stored on a secure server in "
        "Nairobi. It is used only to train the chatbot and is deleted after 3 years.",
    ),
    ("h", "What happens if you say no"),
    ("p", "Nothing changes. You get the same help, and your call is not used."),
    ("h", "Calls we never use"),
    ("p", "Calls about violence, child protection, medical emergencies or health conditions."),
    ("h", "Changing your mind"),
    (
        "p",
        "Send STOP by SMS to 5050 before 30 June 2026, or tell any helpline worker. "
        "Contact the data protection lead, Laura Bianchi, at laura.bianchi@example.org.",
    ),
    ("h", "Record of consent (filled in by the helpline worker)"),
    ("p", "Caller name: Hassan Abdi. Phone: +254 733 555 621. Case ID: KK-2026-00418."),
    ("p", "Language: Somali. Consent given: yes, verbally, on 18 March 2026 at 10:42."),
    ("p", "Helpline worker: Esther Nakiru, staff ID NRN-1187."),
]

POLICY_PDF = [
    ("title", "Chatbot pilot: data protection policy"),
    ("note", NOTE),
    ("p", "Northbridge Relief Network (fictional). Approved by the board on 2 April 2026."),
    ("h", "1. Scope"),
    (
        "p",
        "This policy covers all data used to build, train, test and run the helpline chatbot "
        "piloted in Kakuma and Kalobeyei, Kenya.",
    ),
    ("h", "2. Principles"),
    (
        "p",
        "Do no harm: no data is used if its use could expose a person to risk. Humanity: the "
        "chatbot exists to reduce suffering, not to collect data. Impartiality: the chatbot must "
        "work equally well in Somali, Swahili, Turkana and English. Independence: no data is "
        "shared with authorities, donors or technology companies.",
    ),
    ("h", "3. Data we use"),
    (
        "p",
        "Phase 1: synthetic conversations written by staff and public pages of partner NGOs. "
        "Phase 2, only after re-consent and a re-identification test: anonymised helpline "
        "transcripts. Never: calls about gender-based violence, child protection or health.",
    ),
    ("h", "4. Anonymisation"),
    (
        "p",
        "Names, phone numbers, emails, case IDs, medical record numbers and locations below "
        "district level are removed before data leaves the country office. A re-identification "
        "test on 500 transcripts must find no person before Phase 2 starts.",
    ),
    ("h", "5. Language model provider"),
    (
        "p",
        "The model is fine-tuned on servers we control. If an external API is used, only "
        "sanitized text is sent, and the provider must not train on our data.",
    ),
    ("h", "6. Accountability"),
    (
        "p",
        "The data protection lead reviews this policy every 12 months. Complaints go to "
        "privacy@example.org. A community feedback session is held every 6 months.",
    ),
]

INVENTORY = [
    ("Dataset", "Source", "Contains personal data", "Consent basis", "Status", "Owner"),
    ("Helpline transcripts 2023-2025", "Helpline", "yes", "re-consent needed", "on hold", "Tomás Herrera"),
    ("Synthetic conversations", "Staff-written", "no", "not needed", "in use", "Tomás Herrera"),
    ("Partner NGO web pages", "Public web", "no", "licence check", "in use", "Peter Schmid"),
    ("Clinic referral notes", "Kakuma clinic", "yes (health)", "none", "excluded", "Dr. Samuel Otieno"),
    ("Community FAQ feedback", "Community sessions", "yes", "verbal consent", "in review", "Amina Wanjiru"),
]  # fmt: skip


def make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    for kind, text in CONSENT_FORM:
        if kind == "title":
            doc.add_heading(text, level=0)
        elif kind == "h":
            doc.add_heading(text, level=2)
        elif kind == "note":
            doc.add_paragraph().add_run(text).italic = True
        else:
            doc.add_paragraph(text)
    doc.save(path)


def make_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    y = 60
    for kind, text in POLICY_PDF:
        size = {"title": 16, "h": 12}.get(kind, 10)
        font = "helv" if kind == "p" else ("hebo" if kind in ("title", "h") else "heit")
        box = pymupdf.Rect(60, y, 535, y + 200)
        left = page.insert_textbox(box, text, fontsize=size, fontname=font)
        used = 200 - left if left >= 0 else 200
        y += used + (14 if kind in ("title", "h") else 8)
        if y > 760:
            page, y = doc.new_page(), 60
    doc.save(path)


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data inventory"
    ws.append([NOTE])
    for row in INVENTORY:
        ws.append(list(row))
    wb.save(path)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    make_docx(OUT / "consent_form_template.docx")
    make_pdf(OUT / "chatbot_pilot_data_policy.pdf")
    make_xlsx(OUT / "data_inventory.xlsx")
    print(f"Wrote 3 files to {OUT}")
